"""Load Excel file and map columns to standard format."""
from pathlib import Path

from openpyxl import load_workbook
import pandas as pd
from backend.config import COL_INDICES, EXCLUDE_INCIDENT_TYPES, EXCLUDE_OUTCOMES
from backend.pipeline.anonymize import anonymize_text


class TableFormatError(ValueError):
    """Загруженный файл не соответствует ожидаемому формату реестра обращений."""


# Колонки, без которых обработка не имеет смысла: их отсутствие означает,
# что загружен файл другой структуры.
REQUIRED_COLUMNS = {
    "incident_text": "«Текст инцидента»",
    "municipality": "«Муниципалитет»",
}


def _worksheet_row_count(filepath: str) -> int | None:
    """Count source worksheet data rows, excluding the header row."""
    if Path(filepath).suffix.lower() not in {".xlsx", ".xlsm"}:
        return None

    try:
        wb = load_workbook(filepath, read_only=True, data_only=True)
    except Exception:
        return None

    try:
        return max((wb.active.max_row or 1) - 1, 0)
    finally:
        wb.close()


def _value_counts(frame: pd.DataFrame, column: str) -> dict:
    """Counts per value of a column, with empty values bucketed as «Не указано»."""
    if column not in frame.columns or frame.empty:
        return {}
    series = frame[column].fillna("").astype(str).str.strip().replace("", "Не указано")
    return {str(key): int(value) for key, value in series.value_counts().items()}


def _month_counts(frame: pd.DataFrame, column: str = "date_closed") -> dict:
    """Counts per month bucket of a datetime column (ISO month start -> count)."""
    if column not in frame.columns or frame.empty:
        return {}
    dates = frame[column].dropna()
    if dates.empty:
        return {}
    buckets = dates.dt.to_period("M").dt.to_timestamp()
    return {
        timestamp.isoformat(): int(count)
        for timestamp, count in buckets.value_counts().sort_index().items()
    }


def _resolution_stats(frame: pd.DataFrame) -> dict:
    """Sum and count of valid resolution durations (in days) for closed rows.

    Нужно, чтобы средний срок решения учитывал и «закрытые до анализа» строки:
    из них самих хранится только агрегат (сумма дней и количество корректных
    интервалов), а средний по всем решённым считается на дашборде.
    """
    if (
        "date_created" not in frame.columns
        or "date_closed" not in frame.columns
        or frame.empty
    ):
        return {"sum_days": 0.0, "count": 0}
    created = frame["date_created"]
    closed = frame["date_closed"]
    valid = created.notna() & closed.notna() & (closed >= created)
    if not bool(valid.any()):
        return {"sum_days": 0.0, "count": 0}
    days = (closed[valid] - created[valid]).dt.total_seconds() / 86400.0
    return {"sum_days": float(days.sum()), "count": int(valid.sum())}


def load_excel(filepath: str) -> tuple[pd.DataFrame, dict]:
    """Load Excel, extract relevant columns, and apply status filters.

    Returns (filtered_df, stats) where stats has counts:
        raw_count, dropped_short_text, dropped_incident_type, dropped_outcome,
        and "prefiltered" — aggregated breakdowns of the dropped rows
        (закрытые до анализа и нерешаемые) for the resolved-analytics page.
    """
    selected_indices = sorted(set(COL_INDICES.values()))
    index_to_name = {idx: name for name, idx in COL_INDICES.items()}

    df = None
    try:
        for engine in ("calamine", "openpyxl"):
            try:
                df = pd.read_excel(
                    filepath,
                    header=0,
                    dtype=str,
                    engine=engine,
                    usecols=selected_indices,
                )
                rename_map = {}
                for offset, original_idx in enumerate(selected_indices):
                    if offset < len(df.columns) and original_idx in index_to_name:
                        rename_map[df.columns[offset]] = index_to_name[original_idx]
                df = df.rename(columns=rename_map)
                break
            except (ValueError, ImportError):
                if engine == "openpyxl":
                    df = pd.read_excel(filepath, header=0, dtype=str, engine="openpyxl")
                    columns = df.columns.tolist()
                    col_map = {}
                    for name, idx in COL_INDICES.items():
                        if idx < len(columns):
                            col_map[columns[idx]] = name
                    relevant_cols = [columns[idx] for idx in COL_INDICES.values() if idx < len(columns)]
                    df = df[relevant_cols].rename(columns=col_map)
    except TableFormatError:
        raise
    except Exception as exc:
        raise TableFormatError(
            "Несоответствие формату таблицы: файл не читается как Excel-реестр обращений."
        ) from exc

    if df is None:
        raise TableFormatError("Несоответствие формату таблицы: файл не читается как Excel-реестр обращений.")

    # Файл другой структуры (меньше колонок, другой порядок) — говорим явно,
    # а не падаем позже с KeyError при обработке.
    missing = [label for column, label in REQUIRED_COLUMNS.items() if column not in df.columns]
    if missing:
        raise TableFormatError(
            "Несоответствие формату таблицы: в файле не найдены колонки "
            + ", ".join(missing)
            + ". Загрузите исходный реестр обращений с тем же набором столбцов."
        )

    raw_count = _worksheet_row_count(filepath) or len(df)

    # Parse dates (для timeline-аналитики)
    for date_col in ("date_created", "date_closed"):
        if date_col in df.columns:
            df[date_col] = pd.to_datetime(df[date_col], errors="coerce", dayfirst=True)

    # Clean text
    if "incident_text" in df.columns:
        df["incident_text"] = df["incident_text"].fillna("").astype(str).str.strip()
        df["incident_text"] = df["incident_text"].str.replace(r"<[^>]+>", " ", regex=True)
        # Обезличивание (упоминания, телефоны, email) + маскирование мата.
        df["incident_text"] = df["incident_text"].map(anonymize_text)
        df["incident_text"] = df["incident_text"].str.replace(r"\s+", " ", regex=True).str.strip()

    # Track drops separately
    before_short = len(df)
    df = df[df["incident_text"].str.len() > 10]
    dropped_short_text = before_short - len(df)

    # Колонки на месте, но текстов нет вообще — почти наверняка в файле
    # другая структура (текст лежит не в том столбце).
    if raw_count > 0 and len(df) == 0:
        raise TableFormatError(
            "Несоответствие формату таблицы: колонка с текстом обращений пуста. "
            "Проверьте, что структура файла совпадает с исходным реестром."
        )

    dropped_incident_type = 0
    unsolvable_df = df.iloc[0:0]
    if "incident_type" in df.columns:
        unsolvable_mask = df["incident_type"].fillna("").isin(EXCLUDE_INCIDENT_TYPES)
        unsolvable_df = df[unsolvable_mask]
        df = df[~unsolvable_mask]
        dropped_incident_type = len(unsolvable_df)

    dropped_outcome = 0
    resolved_df = df.iloc[0:0]
    if "outcome" in df.columns:
        resolved_mask = df["outcome"].notna() & df["outcome"].isin(EXCLUDE_OUTCOMES)
        resolved_df = df[resolved_mask]
        df = df[~resolved_mask]
        dropped_outcome = len(resolved_df)

    stats = {
        "raw_count": raw_count,
        "dropped_short_text": dropped_short_text,
        "dropped_incident_type": dropped_incident_type,
        "dropped_outcome": dropped_outcome,
        "prefiltered": {
            "resolved": {
                "total": int(dropped_outcome),
                "by_municipality": _value_counts(resolved_df, "municipality"),
                "by_category": _value_counts(resolved_df, "group"),
                "by_outcome": _value_counts(resolved_df, "outcome"),
                "timeline": _month_counts(resolved_df),
                "resolution": _resolution_stats(resolved_df),
            },
            "unsolvable": {
                "total": int(dropped_incident_type),
                "by_municipality": _value_counts(unsolvable_df, "municipality"),
                "by_category": _value_counts(unsolvable_df, "group"),
            },
        },
    }
    return df.reset_index(drop=True), stats
