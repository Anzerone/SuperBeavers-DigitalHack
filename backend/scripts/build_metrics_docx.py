"""Сборка .docx-отчёта по метрикам классификатора обращений."""
from __future__ import annotations

import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Inches

OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "output", "doc")
OUT_PATH = os.path.abspath(os.path.join(OUT_DIR, "Отчёт_метрики.docx"))

TEAL = RGBColor(0x0D, 0x73, 0x77)
DARK = RGBColor(0x1F, 0x2A, 0x37)
GRAY = RGBColor(0x6B, 0x72, 0x80)
HEADER_BG = "0D7377"
ROW_ALT = "F1F5F9"


def set_cell_bg(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def style_cell_text(cell, *, bold=False, color=None, size=10, align=None, white=False):
    for p in cell.paragraphs:
        if align is not None:
            p.alignment = align
        for run in p.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            if white:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            elif color is not None:
                run.font.color.rgb = color


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        set_cell_bg(hdr[i], HEADER_BG)
        align = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
        style_cell_text(hdr[i], bold=True, white=True, size=10, align=align)
    for r_i, row in enumerate(rows):
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            align = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            style_cell_text(cells[i], bold=(i == 0), size=10, align=align,
                            color=DARK)
            if r_i % 2 == 1:
                set_cell_bg(cells[i], ROW_ALT)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
    return table


def h1(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = TEAL
    p.space_before = Pt(14)
    p.space_after = Pt(6)
    return p


def caption(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(9)
    run.font.italic = True
    run.font.color.rgb = GRAY
    p.space_after = Pt(8)
    return p


def body(doc, text, *, bullet=False):
    p = doc.add_paragraph(style="List Bullet" if bullet else None)
    run = p.add_run(text)
    run.font.size = Pt(10.5)
    run.font.color.rgb = DARK
    return p


def main() -> None:
    os.makedirs(OUT_DIR, exist_ok=True)
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # ---- Титул ----
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Классификатор обращений граждан")
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = TEAL

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Отчёт о метриках качества и производительности")
    r.font.size = Pt(13)
    r.font.color.rgb = GRAY

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run("Омская область  |  Датасет: 419 054 обращения  |  GPU: NVIDIA RTX 5090")
    r.font.size = Pt(10)
    r.font.color.rgb = GRAY

    # ---- 1. Резюме ----
    h1(doc, "1. Краткое резюме")
    body(doc, "Система обрабатывает 419 054 обращения, отбирает проблемные, определяет "
              "тяжесть и категорию, группирует повторяющиеся проблемы по районам и строит "
              "аналитический дашборд. Ниже приведены ключевые показатели качества и скорости.")
    add_table(doc,
        ["Показатель", "Значение"],
        [
            ["Обработано обращений", "47 153 из 419 054 (11,3%)"],
            ["Выявлено проблем", "42 294 (89,7%)"],
            ["Категория по тексту (top-1 / top-3)", "74,9% / 93,1%"],
            ["Качество эмбеддингов (1-NN)", "75,3%"],
            ["Определение проблемы (F1 vs LLM)", "0,965"],
            ["Тяжесть (в пределах +/-1 уровня)", "97,3%"],
            ["Кластеризация: сжатие / шум", "125,9x / 1,7%"],
            ["Полный прогон на GPU", "~5-15 мин"],
        ],
        col_widths=[3.6, 3.0])

    # ---- 2. Данные и покрытие ----
    h1(doc, "2. Данные и покрытие")
    body(doc, "Из исходного файла отсеиваются закрытые, решённые, перенаправленные и "
              "нерешаемые обращения - они не являются актуальными проблемами. Эмбеддинги "
              "считаются только для оставшихся записей.")
    add_table(doc,
        ["Метрика", "Значение", "Доля"],
        [
            ["Строк в исходном файле", "419 054", "100%"],
            ["Прошло фильтры (обработано)", "47 153", "11,3%"],
            ["Отфильтровано", "371 901", "88,7%"],
            ["Дубли текста (дедуп)", "1 020", "2,2%"],
            ["Уникальных текстов", "46 133", "97,8%"],
        ],
        col_widths=[3.4, 1.8, 1.4])

    # ---- 3. Продуктовые метрики ----
    h1(doc, "3. Продуктовые метрики")
    body(doc, "Распределение проблемных обращений по тяжести:")
    add_table(doc,
        ["Тяжесть", "Кол-во", "Доля"],
        [
            ["CRITICAL", "2 103", "5,0%"],
            ["HIGH", "16 523", "39,1%"],
            ["MEDIUM", "14 676", "34,7%"],
            ["LOW", "8 992", "21,3%"],
        ],
        col_widths=[2.6, 2.0, 2.0])

    body(doc, "Распределение по тональности (проблемные):")
    add_table(doc,
        ["Тональность", "Кол-во", "Доля"],
        [
            ["Нейтральная", "36 680", "86,7%"],
            ["Гнев", "4 413", "10,4%"],
            ["Спокойная", "864", "2,0%"],
            ["Отчаяние", "337", "0,8%"],
        ],
        col_widths=[2.6, 2.0, 2.0])

    body(doc, "Топ категорий и муниципалитетов по числу проблем:")
    add_table(doc,
        ["Категория", "Проблем", "Муниципалитет", "Проблем"],
        [
            ["Дороги", "14 969", "Омск г.о.", "32 007"],
            ["ЖКХ", "8 808", "Омская обл., другое", "2 384"],
            ["Благоустройство", "6 947", "Омский район", "1 365"],
            ["Общественный транспорт", "3 254", "Калачинский район", "810"],
            ["Обращение с отходами", "1 946", "Исилькульский район", "787"],
            ["Здравоохранение", "1 845", "Любинский район", "599"],
        ],
        col_widths=[2.4, 1.2, 2.2, 1.0])

    # ---- 4. Технические метрики ----
    h1(doc, "4. Технические метрики")

    body(doc, "4.1. Классификация категории (эталон - человеческая разметка «Группа тем», 26 классов). "
              "Главная честная метрика - предсказание категории по тексту обращения:")
    add_table(doc,
        ["Метрика", "По тексту (честно)", "С префиксом группы*"],
        [
            ["Accuracy top-1", "74,9%", "97,5%"],
            ["Accuracy top-3", "93,1%", "99,9%"],
            ["Macro-F1", "0,519", "0,919"],
            ["Weighted-F1", "0,755", "0,975"],
        ],
        col_widths=[2.6, 2.2, 2.2])
    caption(doc, "* Вариант с префиксом группы во входе содержит утечку эталона (категория "
                 "зашита во вход) и приведён для прозрачности. В рабочем режиме категория "
                 "копируется из готовой «Группы тем», когда она задана; модель используется, "
                 "когда группа отсутствует.")

    body(doc, "4.2. Качество эмбеддингов (bge-m3, чистый текст), эталон - «Группа тем»:")
    add_table(doc,
        ["Метрика", "Значение"],
        [
            ["1-NN accuracy (ближайший сосед той же категории)", "75,3%"],
            ["5-NN purity (доля соседей той же категории)", "72,9%"],
        ],
        col_widths=[4.6, 2.0])

    body(doc, "4.3. Кластеризация (DBSCAN по «район x категория», eps=0.35):")
    add_table(doc,
        ["Метрика", "Значение"],
        [
            ["Кластеров найдено", "336"],
            ["Шум (вне кластеров)", "708 (1,7%)"],
            ["Сжатие (обращений на кластер)", "125,9x"],
            ["Размер кластера (медиана / макс)", "7 / 10 991"],
            ["Silhouette (по категориям)", "0,093"],
        ],
        col_widths=[4.6, 2.0])
    caption(doc, "Низкий silhouette закономерен: тексты жалоб семантически пересекаются между "
                 "категориями (ЖКХ, Благоустройство и Дороги частично смешиваются).")

    body(doc, "4.4. Валидация определения проблемы и тяжести на выборке из 120 обращений, "
              "размеченных независимым LLM-экспертом (qwen3:8b):")
    add_table(doc,
        ["Метрика (is_problem)", "Значение"],
        [
            ["Accuracy", "93,3%"],
            ["Precision", "99,1%"],
            ["Recall", "94,0%"],
            ["F1", "0,965"],
            ["Матрица (TP / FP / FN / TN)", "110 / 1 / 7 / 2"],
        ],
        col_widths=[4.0, 2.6])
    add_table(doc,
        ["Метрика (severity)", "Значение"],
        [
            ["Точное совпадение", "50,0%"],
            ["В пределах +/-1 уровня", "97,3%"],
        ],
        col_widths=[4.0, 2.6])

    # ---- 5. Оговорки ----
    h1(doc, "5. Оговорки и корректность измерений")
    body(doc, "Категория 74,9% - честная способность модели по тексту; значение 97,5% завышено "
              "утечкой и в продукте не показатель, т.к. категория берётся из готовой «Группы тем».", bullet=True)
    body(doc, "Метрика is_problem F1 0,965 измерена относительно LLM-эксперта (не человека). "
              "База перекошена: 89,7% записей - проблемы, поэтому негативный класс мал.", bullet=True)
    body(doc, "Тяжесть субъективна: точное совпадение 50%, но грубых ошибок почти нет - 97,3% "
              "предсказаний в пределах одного уровня.", bullet=True)
    body(doc, "Macro-F1 категории (0,519) ограничен дисбалансом: редкие категории "
              "предсказываются хуже частых (Дороги, ЖКХ).", bullet=True)

    # ---- 6. Производительность ----
    h1(doc, "6. Производительность")
    body(doc, "После фильтрации обрабатывается ~46 тыс. текстов (не 419 тыс.). Узкое место - "
              "расчёт эмбеддингов; на GPU оно снимается.")
    add_table(doc,
        ["Сценарий", "CUDA (RTX 5090)", "CPU"],
        [
            ["Эмбеддинги 46k", "~1 мин", "~128 мин"],
            ["Полный прогон файла", "~5-15 мин", "~2,5 часа"],
            ["Fine-tuning (LoRA, 1 эпоха)", "~18 мин (26k)", "неприменимо"],
        ],
        col_widths=[2.8, 2.2, 1.8])

    # ---- 7. Методология ----
    h1(doc, "7. Методология")
    body(doc, "Эмбеддинги: BAAI/bge-m3, max_seq_length=320, fp16 на GPU.", bullet=True)
    body(doc, "Классификатор: Logistic Regression (One-vs-Rest) поверх эмбеддингов.", bullet=True)
    body(doc, "Кластеризация: DBSCAN (cosine, eps=0.35) отдельно по каждой паре «район x категория».", bullet=True)
    body(doc, "Эталон категории: колонка «Группа тем»; train/test split 70/30.", bullet=True)
    body(doc, "Эталон is_problem/severity: независимая разметка LLM qwen3:8b на 120 обращениях.", bullet=True)
    body(doc, "Все измерения выполнены на реальном датасете 419 054 обращений.", bullet=True)

    doc.save(OUT_PATH)
    print(f"Сохранено: {OUT_PATH}")


if __name__ == "__main__":
    main()
